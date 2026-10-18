"""Transfer-matrix model for multilayer optical directional filters (SI units)."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Literal

import numpy as np
from openpyxl import load_workbook


POLARIZATION = Literal["s", "p"]
LOW_INDEX = Literal["CaF2", "SiO2", "Al2O3"]

ETA_0 = 376.7

TIO2_NUMBER = 1
SI_NUMBER = 2
ZNS_NUMBER = 3
VO2_NUMBER = 4
WEYL_NUMBER = 5


@dataclass(frozen=True)
class MaterialData:
    wavelengths: np.ndarray
    index: np.ndarray


@dataclass(frozen=True)
class FilterResult:
    wavelength: np.ndarray
    theta_deg: np.ndarray
    s11: np.ndarray
    s21: np.ndarray
    reflectance: np.ndarray
    transmission: np.ndarray
    absorption: np.ndarray
    absorption_matrix: np.ndarray
    transmission_matrix: np.ndarray
    reflectance_matrix: np.ndarray


def _read_range(workbook: Path, cell_range: str, sheet: str = "Sheet1") -> np.ndarray:
    wb = load_workbook(workbook, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        values = []
        for row in ws[cell_range]:
            for cell in row:
                values.append(np.nan if cell.value is None else float(cell.value))
    finally:
        wb.close()
    return np.asarray(values, dtype=float)


def _complex_index(
    workbook: Path,
    wavelength_range: str,
    real_range: str,
    imag_range: str | None = None,
) -> MaterialData:
    wavelengths = _read_range(workbook, wavelength_range) * 1e-6
    real = _read_range(workbook, real_range)
    if imag_range is None:
        index = real.astype(complex)
    else:
        index = real + 1j * _read_range(workbook, imag_range)
    return MaterialData(wavelengths=wavelengths, index=index)


def load_material_data(base_dir: str | Path = ".") -> dict[str, MaterialData]:
    """Load the refractive-index tables used by the MATLAB script."""
    base = Path(base_dir)
    refractive_index = base / "RefractiveIndex.xlsx"
    refractive_index_24 = base / "RefractiveIndex24.xlsx"
    missing = [str(path) for path in (refractive_index, refractive_index_24) if not path.is_file()]
    if missing:
        raise FileNotFoundError(
            "Required material input files are missing: " + ", ".join(missing)
            + ". Obtain the original workbooks from the maintainer; see README.md."
        )

    return {
        "TiO2": _complex_index(refractive_index_24, "A3:A850", "B3:B850", "C3:C850"),
        "Si": _complex_index(refractive_index_24, "F3:F142", "G3:G142", "H3:H142"),
        "ZnS": _complex_index(refractive_index, "AD3:AD314", "AE3:AE314", "AF3:AF314"),
        "SiO2": _complex_index(refractive_index_24, "O3:O1912", "P3:P1912", "Q3:Q1912"),
        "VO2ins": _complex_index(refractive_index, "S3:S130", "T3:T130", "U3:U130"),
        "VO2met": _complex_index(refractive_index, "V3:V114", "W3:W114", "X3:X114"),
        "Ta2O5": _complex_index(refractive_index, "AL3:AL214", "AM3:AM214"),
        "CaF2": _complex_index(refractive_index, "AH3:AH1514", "AI3:AI1514", "AJ3:AJ1514"),
        "Al2O3": _complex_index(refractive_index, "AP3:AP1514", "AQ3:AQ1514", "AR3:AR1514"),
        "Weyl": _complex_index(refractive_index, "AX3:AX65", "AY3:AY65", "AZ3:AZ65"),
    }


def matlab_colon(start: float, step: float, stop: float) -> np.ndarray:
    """Match MATLAB's start:step:stop behavior for the grids used here."""
    count = int(np.floor((stop - start) / step + 1e-12)) + 1
    return start + step * np.arange(count, dtype=float)


def interp1_with_end_fill(x: np.ndarray, y: np.ndarray, query: float) -> complex:
    """Linear interp1 equivalent with MATLAB's explicit y(end) fill value."""
    x = np.asarray(x, dtype=float)
    y = np.asarray(y, dtype=complex)
    fill = y[-1]

    if query < x[0] or query > x[-1] or np.isnan(query):
        return complex(fill)

    real = np.interp(query, x, y.real)
    imag = np.interp(query, x, y.imag)
    return complex(real, imag)


def _permittivity(n: complex) -> complex:
    return n.real**2 - n.imag**2 + 1j * 2 * n.real * n.imag


def _default_thicknesses() -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    hthick_1 = [40e-9, 70e-9, 85e-9, 70e-9, 95e-9, 95e-9, 230e-9, 0, 0]
    lthick_1 = [80e-9, 100e-9, 125e-9, 100e-9, 125e-9, 135e-9, 485e-9, 0, 0]
    hthick_2 = [450e-9, 135e-9, 180e-9, 162.5e-9, 615.1e-9, 275.6e-9, 162.5e-9, 615.1e-9, 275.6e-9]
    lthick_2 = [2550e-9, 175e-9, 435e-9, 1019.2e-9, 448.9e-9, 0, 1019.2e-9, 448.9e-9, 0]
    hthick_3 = [120e-9, 240e-9, 0, 0, 0, 0, 0, 0, 0]
    lthick_3 = [160e-9, 700e-9, 0, 0, 0, 0, 0, 0, 0]
    hthick_4 = [450e-9, 0, 0, 0, 0, 0, 0, 0, 0]
    lthick_4 = [2550e-9, 0, 0, 0, 0, 0, 0, 0, 0]
    hthick_5 = [450e-9, 70e-9, 80e-9, 100e-9, 115e-9, 135e-9, 0, 0, 0]
    lthick_5 = [2550e-9, 93.1e-9, 110e-9, 121.36e-9, 132.89e-9, 165.44e-9, 0, 0, 0]

    n_pairs_1 = [0, 0, 0, 0, 0, 0, 0, 0, 0]
    n_pairs_2 = [5, 0, 0, 0, 0, 0, 0, 0, 0]
    n_pairs_3 = [0, 0, 0, 0, 0, 0, 0, 0, 0]
    n_pairs_4 = [0, 0, 0, 0, 0, 0, 0, 0, 0]
    n_pairs_5 = [0, 0, 0, 0, 0, 0, 0, 0, 0]

    hthicks = np.asarray([hthick_1, hthick_2, hthick_3, hthick_4, hthick_5], dtype=float)
    lthicks = np.asarray([lthick_1, lthick_2, lthick_3, lthick_4, lthick_5], dtype=float)
    pair_numbers = np.asarray([n_pairs_1, n_pairs_2, n_pairs_3, n_pairs_4, n_pairs_5], dtype=int)
    return hthicks, lthicks, pair_numbers


def _high_index(
    material: int,
    lam: float,
    material_data: dict[str, MaterialData],
    vo2_insulating: bool,
) -> complex:
    if material == TIO2_NUMBER:
        data = material_data["TiO2"]
    elif material == SI_NUMBER:
        data = material_data["Si"]
    elif material == ZNS_NUMBER:
        data = material_data["ZnS"]
    elif material == VO2_NUMBER:
        data = material_data["VO2ins" if vo2_insulating else "VO2met"]
    elif material == WEYL_NUMBER:
        data = material_data["Weyl"]
    else:
        raise ValueError(f"Unsupported high-index material number: {material}")
    return interp1_with_end_fill(data.wavelengths, data.index, lam)


def _low_index(lam: float, material_data: dict[str, MaterialData], low_index: LOW_INDEX) -> complex:
    data = material_data[low_index]
    return interp1_with_end_fill(data.wavelengths, data.index, lam)


def compute_filter(
    base_dir: str | Path = ".",
    wavelength: np.ndarray | None = None,
    theta_deg: np.ndarray | None = None,
    low_index: LOW_INDEX = "CaF2",
    polarization: POLARIZATION = "s",
    n_materials: int = 2,
    vo2_insulating: bool = True,
    material_data: dict[str, MaterialData] | None = None,
    hthicks: np.ndarray | None = None,
    lthicks: np.ndarray | None = None,
    pair_numbers: np.ndarray | None = None,
) -> FilterResult:
    """Compute the multilayer directional-filter response from the MATLAB script."""
    if polarization not in {"s", "p"}:
        raise ValueError("polarization must be 's' or 'p'")
    if low_index not in {"CaF2", "SiO2", "Al2O3"}:
        raise ValueError("low_index must be 'CaF2', 'SiO2', or 'Al2O3'")
    if not isinstance(n_materials, (int, np.integer)) or not 1 <= n_materials <= 5:
        raise ValueError("n_materials must be an integer between 1 and 5")

    if wavelength is None:
        wavelength = matlab_colon(6e-6, 10e-9, 12e-6)
    else:
        wavelength = np.asarray(wavelength, dtype=float)

    if theta_deg is None:
        theta_deg = matlab_colon(0, 1, 85)
    else:
        theta_deg = np.asarray(theta_deg, dtype=float)

    for name, values in (("wavelength", wavelength), ("theta_deg", theta_deg)):
        if values.ndim != 1 or values.size == 0 or not np.all(np.isfinite(values)):
            raise ValueError(f"{name} must be a nonempty, finite one-dimensional array")
    if np.any(wavelength <= 0):
        raise ValueError("wavelength must contain positive values in meters")
    if np.any(np.abs(theta_deg) >= 90):
        raise ValueError("theta_deg must be strictly between -90 and 90 degrees")

    default_hthicks, default_lthicks, default_pair_numbers = _default_thicknesses()
    if hthicks is None:
        hthicks = default_hthicks
    else:
        hthicks = np.asarray(hthicks, dtype=float)

    if lthicks is None:
        lthicks = default_lthicks
    else:
        lthicks = np.asarray(lthicks, dtype=float)

    if pair_numbers is None:
        pair_numbers = default_pair_numbers
    else:
        raw_pairs = np.asarray(pair_numbers, dtype=float)
        if not np.all(np.isfinite(raw_pairs)) or np.any(raw_pairs < 0) or np.any(raw_pairs != np.floor(raw_pairs)):
            raise ValueError("pair_numbers must contain nonnegative integers")
        pair_numbers = np.asarray(pair_numbers, dtype=int)

    if hthicks.shape != default_hthicks.shape:
        raise ValueError(f"hthicks must have shape {default_hthicks.shape}")
    if lthicks.shape != default_lthicks.shape:
        raise ValueError(f"lthicks must have shape {default_lthicks.shape}")
    if pair_numbers.shape != default_pair_numbers.shape:
        raise ValueError(f"pair_numbers must have shape {default_pair_numbers.shape}")
    for name, values in (("hthicks", hthicks), ("lthicks", lthicks)):
        if not np.all(np.isfinite(values)) or np.any(values < 0):
            raise ValueError(f"{name} must contain finite, nonnegative thicknesses in meters")

    if material_data is None:
        material_data = load_material_data(base_dir)

    theta_rad = theta_deg * np.pi / 180
    k0 = 2 * np.pi / wavelength

    n_columns = len(wavelength) * len(theta_deg)
    s11 = np.zeros(n_columns, dtype=complex)
    s21 = np.zeros(n_columns, dtype=complex)

    for angle_index, theta in enumerate(theta_rad):
        sin2 = np.sin(theta) ** 2
        for q, lam in enumerate(wavelength):
            abcd_final = np.eye(2, dtype=complex)
            z0 = None

            for material in range(1, n_materials + 1):
                nh = _high_index(material, lam, material_data, vo2_insulating)
                nl = _low_index(lam, material_data, low_index)

                eps_h = _permittivity(nh)
                eps_l = _permittivity(nl)
                kh = k0[q] * np.sqrt(eps_h - sin2)
                kl = k0[q] * np.sqrt(eps_l - sin2)

                if polarization == "p":
                    zh = ETA_0 * np.sqrt(eps_h - sin2) / eps_h
                    zl = ETA_0 * np.sqrt(eps_l - sin2) / eps_l
                    z0 = ETA_0 * np.cos(theta)
                else:
                    zh = ETA_0 / np.sqrt(eps_h - sin2)
                    zl = ETA_0 / np.sqrt(eps_l - sin2)
                    z0 = ETA_0 / np.cos(theta)

                for th, tl, n_pairs in zip(hthicks[material - 1], lthicks[material - 1], pair_numbers[material - 1]):
                    abcd_high = np.asarray(
                        [
                            [np.cos(kh * th), 1j * zh * np.sin(kh * th)],
                            [1j * np.sin(kh * th) / zh, np.cos(kh * th)],
                        ],
                        dtype=complex,
                    )
                    abcd_low = np.asarray(
                        [
                            [np.cos(kl * tl), 1j * zl * np.sin(kl * tl)],
                            [1j * np.sin(kl * tl) / zl, np.cos(kl * tl)],
                        ],
                        dtype=complex,
                    )
                    for _ in range(n_pairs):
                        abcd_final = abcd_final @ abcd_high @ abcd_low

            if z0 is None:
                raise ValueError("n_materials must be at least 1")

            a, b = abcd_final[0]
            c, d = abcd_final[1]
            column = q + len(wavelength) * angle_index
            s11[column] = (a * z0 + b - c * z0 * z0 - d * z0) / (a * z0 + b + c * z0 * z0 + d * z0)
            s21[column] = 2 / (a + b / z0 + c * z0 + d)

    reflectance = np.abs(s11) ** 2
    transmission = np.abs(s21) ** 2
    absorption = 1 - reflectance - transmission

    matrix_shape = (len(wavelength), len(theta_deg))
    absorption_matrix = absorption.reshape((len(theta_deg), len(wavelength))).T
    transmission_matrix = transmission.reshape((len(theta_deg), len(wavelength))).T
    reflectance_matrix = reflectance.reshape((len(theta_deg), len(wavelength))).T

    return FilterResult(
        wavelength=wavelength,
        theta_deg=theta_deg,
        s11=s11,
        s21=s21,
        reflectance=reflectance,
        transmission=transmission,
        absorption=absorption,
        absorption_matrix=absorption_matrix.reshape(matrix_shape),
        transmission_matrix=transmission_matrix.reshape(matrix_shape),
        reflectance_matrix=reflectance_matrix.reshape(matrix_shape),
    )


if __name__ == "__main__":
    result = compute_filter(Path(__file__).resolve().parent)
    print(f"wavelength points: {len(result.wavelength)}")
    print(f"angle points: {len(result.theta_deg)}")
    print(f"transmission matrix shape: {result.transmission_matrix.shape}")
    print(f"transmission min/max: {result.transmission.min():.12g} / {result.transmission.max():.12g}")
